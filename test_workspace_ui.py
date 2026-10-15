"""Isolated browser regression for layout, inspector saves and navigation."""
import json
import io
import zipfile
import os
from pathlib import Path
import socket
import subprocess
import sys
import time
import urllib.request
import uuid

from playwright.sync_api import expect, sync_playwright
from docx import Document
from PIL import Image


root = Path(__file__).resolve().parent
output = root / '.smoke_tmp' / ('workspace-ui-' + uuid.uuid4().hex[:8])
output.mkdir(parents=True)
word = Document()
word.add_heading('人物：林晚', level=1)
word.add_paragraph('冬季制服参考')
picture = io.BytesIO()
Image.new('RGB', (180, 240), '#bb3344').save(picture, format='PNG')
word.add_picture(io.BytesIO(picture.getvalue()))
word.add_heading('规则：潮汐', level=1)
word.add_paragraph('潮汐规则示意图')
word.add_picture(io.BytesIO(picture.getvalue()))
word.save(output / 'visual-import.docx')
os.environ.update(DB_PATH=str(output / 'test.db'), WRITEHTML_ADMIN_PASSWORD='UiTestPass123!',
                  INSPIRATION_WORKER_ENABLED='false', PYTHONUTF8='1')
import db

db.init_db()
uid = db.verify_user('admin', 'UiTestPass123!')['id']
work = db.create_work(uid, '雾港 UI 验收')
chapter = db.create_chapter(work['id'], uid, '雨夜来信')
db.update_chapter(chapter['id'], uid, None, '林晚抵达港口。\n' * 100, None)
db.create_chapter(work['id'], uid, '灯塔证人')
entity = db.create_entity(work['id'], uid, '林晚', '人物', '档案修复师', '谨慎，善于观察')
card = db.save_production_card(work['id'], uid, {'category': 'rule', 'name': '潮汐规则',
    'summary': '大潮会送来旧信', 'detail': '信中只有已经发生的事'})
db.save_production_card(work['id'], uid, {'category': 'rule',
    'name': '这是一条用于检查左侧树形目录长名称显示不会横向溢出的世界规则', 'summary': '长名称显示验收'})
book = db.create_story_plan_node(work['id'], uid, {'node_type': 'book', 'title': '雾港往事'})
db.create_story_plan_node(work['id'], uid, {'node_type': 'chapter', 'title': '追查来信',
    'chapter_id': chapter['id'], 'parent_id': book['id'], 'goal': '找出寄信人'})
db.save_conversation(uid, chapter['id'], [
    {'role': 'user', 'content': '查一下当前资料'},
    {'role': 'tool', 'name': 'read_chapter', 'content': json.dumps({'summary': '已读正文'})},
    {'role': 'tool', 'name': 'list_characters', 'content': json.dumps({'summary': '已读人物'})},
    {'role': 'tool', 'name': 'save_character_card', 'content': json.dumps({'summary': '已保存人物'})},
    {'role': 'tool', 'name': 'read_chapter', 'content': json.dumps({'error': '测试错误应保持可见'})},
    {'role': 'assistant', 'content': '资料已读取。'},
], '')

with socket.socket() as sock:
    sock.bind(('127.0.0.1', 0))
    port = sock.getsockname()[1]
base = f'http://127.0.0.1:{port}'
log = (output / 'server.log').open('w', encoding='utf-8')
server = subprocess.Popen([sys.executable, '-m', 'uvicorn', 'main:app', '--host', '127.0.0.1',
                           '--port', str(port), '--log-level', 'warning'], cwd=root, stdout=log, stderr=log)
try:
    for _ in range(150):
        try:
            urllib.request.urlopen(base, timeout=.5).close()
            break
        except OSError:
            if server.poll() is not None:
                raise RuntimeError((output / 'server.log').read_text(encoding='utf-8'))
            time.sleep(.1)
    else:
        raise RuntimeError('Local UI server did not start')
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel=os.getenv('BROWSER_CHANNEL') or None)
        context = browser.new_context(viewport={'width': 1440, 'height': 900})
        page = context.new_page()
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(base)
        page.get_by_placeholder('用户名', exact=True).fill('admin')
        page.get_by_placeholder('密码', exact=True).fill('UiTestPass123!')
        page.get_by_role('button', name='登录', exact=True).click()
        expect(page.locator('#chapterSearch')).to_be_visible()
        expect(page.locator('.c-title').filter(has_text='雨夜来信')).to_be_visible()
        page.locator('#chapterSearch').fill('1')
        page.locator('#chapterSearch').press('Enter')
        expect(page.locator('#chapTitle')).to_have_value('雨夜来信')
        page.locator('#chapterSearch').fill('')
        page.locator('#aiToggleBtn').click()
        expect(page.locator('.agent-tool-group')).to_have_count(1)
        expect(page.locator('.tool-danger')).to_be_visible()
        page.locator('.agent-tool-group summary').click()
        expect(page.locator('.agent-tool-group .cm.action')).to_have_count(3)
        for width in (1920, 1440, 1024):
            page.set_viewport_size({'width': width, 'height': 900})
            page.wait_for_function("document.querySelector('.main').getBoundingClientRect().right <= document.querySelector('#aiSide').getBoundingClientRect().left + 1 && document.querySelector('.main').getBoundingClientRect().width >= 540")
            bounds = page.locator('.main').bounding_box()
            assert bounds['width'] >= 540
            assert page.evaluate('document.body.scrollWidth <= innerWidth')
            page.screenshot(path=str(output / f'editor-{width}.png'))
        page.locator('#storyToggleBtn').click()
        page.wait_for_function("document.querySelector('#storyDrawer').getBoundingClientRect().left < 2")
        assert page.locator('#storyDrawer').bounding_box()['y'] >= 50
        assert page.locator('#aiSide').is_visible()
        page.locator('#storyDrawer button[title="关闭"]').click()
        page.set_viewport_size({'width': 1440, 'height': 900})
        page.locator('#content').evaluate('el => { el.setSelectionRange(25, 35); el.scrollTop = 250; }')
        before = page.locator('#content').evaluate('el => ({start:el.selectionStart,end:el.selectionEnd,scroll:el.scrollTop})')
        page.locator('#productionToggleBtn').click()
        page.locator('#productionCanonBtn').click()
        expect(page.locator('.production-world-groups section')).to_have_count(2)
        tree = page.locator('#productionResourceList')
        rules = tree.locator('[data-folder="rule"]')
        expect(tree.locator('.production-resource-tree-folder')).to_have_count(6)
        expect(tree.locator('.production-resource-row')).to_have_count(3)
        assert page.locator('#productionResourceTabs').count() == 0
        assert tree.evaluate('el => el.scrollWidth <= el.clientWidth')
        rules.get_by_role('button', name='收起规则', exact=True).click()
        expect(rules).to_have_attribute('aria-expanded', 'false')
        expect(rules.locator('[role="group"]')).to_be_hidden()
        page.locator('#productionResourceSearch').fill('大潮')
        expect(rules).to_have_attribute('aria-expanded', 'true')
        expect(tree.locator('.production-resource-row')).to_have_count(1)
        page.locator('#productionResourceSearch').fill('')
        expect(rules).to_have_attribute('aria-expanded', 'false')
        rules.locator('.production-resource-folder-select').click()
        expect(rules).to_have_attribute('aria-expanded', 'true')
        expect(page.locator('.production-world-groups section')).to_have_count(1)
        expect(tree.locator('.production-resource-tree-folder')).to_have_count(6)
        page.locator('#productionResourceSearch').fill('林晚')
        expect(tree.locator('[data-folder="character"]')).to_be_visible()
        expect(tree.locator('.production-resource-row')).to_have_count(1)
        page.locator('#productionResourceSearch').fill('不存在的资源')
        expect(tree).to_contain_text('没有匹配的资源')
        page.locator('#productionResourceSearch').fill('')
        tree.locator('[data-tree-key="all"]').click()
        expect(page.locator('.production-world-groups section')).to_have_count(2)
        root_item = tree.locator('[data-tree-key="all"]')
        root_item.focus()
        root_item.press('ArrowDown')
        characters = tree.locator('[data-folder="character"]')
        expect(characters).to_be_focused()
        characters.press('ArrowLeft')
        expect(characters).to_have_attribute('aria-expanded', 'false')
        characters.press('ArrowRight')
        expect(characters).to_have_attribute('aria-expanded', 'true')
        characters.press('ArrowRight')
        expect(tree.locator(f'[data-tree-key="character:{entity["id"]}"]')).to_be_focused()
        tree.locator(f'[data-tree-key="character:{entity["id"]}"]').press('ArrowLeft')
        expect(characters).to_be_focused()
        characters.press('Home')
        expect(root_item).to_be_focused()
        tree.locator('[data-folder="skill"] .production-resource-tree-toggle').click()
        page.screenshot(path=str(output / 'resource-tree-1440.png'), animations='disabled')
        page.locator('#agentDocumentFiles').set_input_files(str(output / 'visual-import.docx'))
        expect(page.locator('.agent-document-chip')).to_contain_text('2 图')
        page.locator('.agent-document-chip').get_by_role('button', name='查看图片', exact=True).click()
        expect(page.locator('#resourceImageGrid .resource-image-tile')).to_have_count(2)
        with db.get_conn() as connection:
            image_id = connection.execute('SELECT id FROM resource_images WHERE source_index=1').fetchone()[0]
        page.locator(f'#resourceImageGrid [data-resource-image-id="{image_id}"]').click()
        page.locator('#resourceImageOwner').select_option(f'character:{entity["id"]}')
        page.locator('#resourceImageLabel').fill('林晚·冬季制服')
        page.locator('#resourceImageVariant').select_option('outfit')
        page.get_by_role('button', name='保存绑定', exact=True).click()
        expect(page.locator('#resourceImageMessage')).to_have_text('已保存')
        page.get_by_role('button', name='设为主图', exact=True).click()
        expect(page.locator('#resourceImageEditor')).not_to_contain_text('设为主图')
        page.locator('#resourceImageEditor').get_by_role('button', name='复制', exact=True).click()
        expect(page.locator('#resourceImageOwner')).to_have_value('')
        page.locator('#resourceImageOwner').select_option(f'character:{entity["id"]}')
        expect(page.locator(f'#resourceImageParent option[value="{image_id}"]')).to_have_count(1)
        page.locator('#resourceImageParent').select_option(str(image_id))
        page.locator('#resourceImageLabel').fill('林晚·夏季制服')
        page.locator('#resourceImageVariant').select_option('outfit')
        page.get_by_role('button', name='保存绑定', exact=True).click()
        expect(page.locator('#resourceImageMessage')).to_have_text('已保存')
        for width in (1440, 1024, 390):
            page.set_viewport_size({'width': width, 'height': 844})
            expect(page.locator('#resourceImageEditor')).to_be_visible()
            assert page.evaluate('document.body.scrollWidth <= innerWidth')
            editor = page.locator('#resourceImageEditor').bounding_box()
            assert editor['x'] >= 0 and editor['x'] + editor['width'] <= width + 1
            page.screenshot(path=str(output / f'visual-gallery-{width}.png'), animations='disabled')
        page.get_by_role('button', name='关闭图片素材库', exact=True).click()
        page.set_viewport_size({'width': 1440, 'height': 900})
        with page.expect_download() as downloaded:
            page.locator('#productionWorkspace button[title="导出创作素材包"]').click()
        downloaded.value.save_as(output / 'creative-package.zip')
        with zipfile.ZipFile(output / 'creative-package.zip') as package:
            manifest = json.loads(package.read('manifest.json'))
            assert len(manifest['images']) == 3
            assert any(image['parent_image_ref'] == f'image:{image_id}' for image in manifest['images'])
        tree.locator(f'[data-tree-key="card:{card["id"]}"]').click()
        expect(page.locator('#productionCardName')).to_be_visible()
        page.locator('#productionInspectorTabs').get_by_role('tab', name='图片', exact=True).click()
        expect(page.locator('[data-inspector-panel="images"]')).to_contain_text('管理图片')
        page.locator('#productionInspectorTabs').get_by_role('tab', name='基础设定', exact=True).click()
        expect(page.locator('#productionCardScope')).to_be_hidden()
        page.locator('#productionCardName').fill('潮汐规则修订')
        page.get_by_role('tab', name='本章状态', exact=True).click()
        page.locator('#productionCardState').fill('状态：已被主角发现')
        expect(page.locator('[data-save-kind="state"]')).to_contain_text('待记录')
        expect(page.locator('[data-save-kind="base"]')).to_contain_text('已同步', timeout=10000)
        assert db.get_production_card(card['id'], uid)['name'] == '潮汐规则修订'
        assert not db.get_production_card(card['id'], uid)['versions']
        page.get_by_role('button', name='记录本章状态', exact=True).click()
        expect(page.locator('[data-save-kind="state"]')).to_contain_text('已记录', timeout=10000)
        page.get_by_role('tab', name='历史版本', exact=True).click()
        expect(page.locator('[data-inspector-panel="history"]')).to_contain_text('已被主角发现')
        page.locator('#productionInspector button[title="关闭详情"]').click()
        page.locator('.production-world-groups button').filter(has_text='潮汐规则修订').click()
        expect(page.locator('[data-save-kind="state"]')).to_contain_text('已记录')
        assert '待记录' not in page.locator('[data-save-kind="state"]').inner_text()
        page.locator('#productionInspector button[title="关闭详情"]').click()
        page.set_viewport_size({'width': 1440, 'height': 900})
        page.reload()
        expect(page.locator('.c-title').filter(has_text='雨夜来信')).to_be_visible()
        page.locator('#chapterSearch').fill('1')
        page.locator('#chapterSearch').press('Enter')
        expect(page.locator('#chapTitle')).to_have_value('雨夜来信')
        page.locator('#chapterSearch').fill('')
        page.locator('#content').evaluate('el => { el.setSelectionRange(25, 35); el.scrollTop = 250; }')
        before = page.locator('#content').evaluate('el => ({start:el.selectionStart,end:el.selectionEnd,scroll:el.scrollTop})')
        page.locator('#productionToggleBtn').click()
        expect(page.locator('#productionResourceList [data-folder="skill"]')).to_have_attribute('aria-expanded', 'false')
        page.locator('.production-world-groups button').filter(has_text='潮汐规则修订').click()
        expect(page.locator('[data-save-kind="state"]')).to_contain_text('已记录')
        page.get_by_role('tab', name='本章状态', exact=True).click()
        assert '已被主角发现' in page.locator('#productionCardState').input_value()
        page.locator('#productionInspector button[title="关闭详情"]').click()
        page.locator('#productionPlanBtn').click()
        page.locator('.story-plan-route-main').filter(has_text='追查来信').click()
        expect(page.locator('#productionBreadcrumb')).to_contain_text('雾港往事')
        expect(page.locator('#productionInspectorFooter')).to_be_visible()
        for width in (1440, 1024, 390):
            page.set_viewport_size({'width': width, 'height': 844})
            if width == 390:
                page.locator('#aiSide button[title="收起"]').click()
            page.wait_for_function("document.querySelector('#productionInspector').getBoundingClientRect().right <= innerWidth + 1")
            assert page.evaluate('document.body.scrollWidth <= innerWidth')
            footer = page.locator('#productionInspectorFooter').bounding_box()
            assert footer['y'] + footer['height'] <= 845
            if width == 390:
                assert page.locator('.production-head').bounding_box()['height'] <= 115
            page.screenshot(path=str(output / f'inspector-{width}.png'))
        page.locator('#productionInspector button[title="关闭详情"]').click()
        page.locator('#productionCanonBtn').click()
        for width in (1920, 1024, 390):
            page.set_viewport_size({'width': width, 'height': 844})
            if page.locator('#productionWorkspace').evaluate("el => el.classList.contains('resources-closed')"):
                page.locator('button[title="设定资源库"]').click()
            expect(tree).to_be_visible()
            assert tree.evaluate('el => el.scrollWidth <= el.clientWidth')
            assert page.evaluate('document.body.scrollWidth <= innerWidth')
            if width == 390:
                assert tree.locator('.production-resource-folder-row').first.bounding_box()['height'] >= 44
            page.screenshot(path=str(output / f'resource-tree-{width}.png'), animations='disabled')
        page.locator('button[title="返回正文编辑器"]').click()
        after = page.locator('#content').evaluate('el => ({start:el.selectionStart,end:el.selectionEnd,scroll:el.scrollTop})')
        assert after['start'] == before['start'] and after['end'] == before['end']
        page.locator('#aiToggleBtn').click()
        expect(page.locator('#writingGuideCard.compact')).to_be_visible()
        composer = page.locator('#agentInput').bounding_box()
        assert composer['width'] >= 350
        assert page.locator('#agentMicBtn').bounding_box()['y'] >= composer['y'] + composer['height']
        assert not errors, errors
        browser.close()
    print('Workspace UI checks passed. Screenshots:', output)
    for regression in ('test_story_plan_ui.py', 'test_writing_packs_ui.py'):
        subprocess.run([sys.executable, regression, base], cwd=root,
                       env={**os.environ, 'BROWSER_CHANNEL': os.getenv('BROWSER_CHANNEL', '')}, check=True)
finally:
    server.terminate()
    server.wait(timeout=15)
    log.close()
