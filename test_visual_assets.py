"""Offline regression for Word imports, visual ownership, variants and handoff ZIPs."""
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import zipfile
import threading
from concurrent.futures import ThreadPoolExecutor

from docx import Document
from PIL import Image

root = Path(tempfile.mkdtemp(prefix='writehtml-visual-'))
os.environ.update(DB_PATH=str(root/'test.db'), ENTITY_IMAGE_STORAGE_DIR=str(root/'images'),
                  WRITEHTML_ADMIN_PASSWORD='VisualTest123!', INSPIRATION_WORKER_ENABLED='false')
from fastapi.testclient import TestClient
import db
import main
import visual_assets
import document_imports

client = TestClient(main.app)
token = client.post('/api/login', json={'username':'admin','password':'VisualTest123!'}).json()['token']
headers = {'Authorization':'Bearer '+token}
uid = main._sessions[token]
wid = db.create_work(uid, '图文导入测试')['id']
other_work = db.create_work(uid, '其他作品')['id']
other_uid = db.create_user('visual-other', 'VisualOther123!')['id']
other_token = client.post('/api/login', json={'username':'visual-other','password':'VisualOther123!'}).json()['token']
other_headers = {'Authorization':'Bearer '+other_token}
entity = db.create_entity(wid, uid, '林晚','人物','修复师','黑发，谨慎')
place = db.save_production_card(wid, uid, {'category':'location','name':'雾港','summary':'古老港口'})
prop = db.save_production_card(wid, uid, {'category':'item','name':'钥匙'})
org = db.save_production_card(wid, uid, {'category':'organization','name':'档案馆'})
cross = db.create_entity(other_work, uid, '其他人物','人物','','')

def png(color):
    output=io.BytesIO();Image.new('RGB',(96,64),color).save(output,format='PNG');return output.getvalue()

document=Document()
document.add_heading('故事大纲',level=1);document.add_paragraph('未来计划：主角进入雾港。')
document.add_heading('人物：林晚',level=1);document.add_paragraph('林晚的冬季制服')
document.add_picture(io.BytesIO(png('red')))
document.add_heading('地点：雾港',level=1);document.add_paragraph('雾港完整的广场')
document.add_picture(io.BytesIO(png('green')))
table=document.add_table(rows=1,cols=2);table.cell(0,0).text='道具';table.cell(0,1).text='古铜钥匙'
document.add_heading('第1章 雨夜',level=1);document.add_paragraph('林晚带着钥匙走进雾港。这是已经写好的正文。')
document.add_heading('第2章 灯塔',level=1);document.add_paragraph('她在灯塔遇见沈舟。')
data=io.BytesIO();document.save(data);word=data.getvalue()
response=client.post('/api/agent/documents/extract',content=word,
    headers={**headers,'X-Work-Id':str(wid),'X-File-Name':'test.docx'})
assert response.status_code==200,response.text
imported=response.json();doc_id=imported['document_id'];images=imported['images']
assert len(images)==2 and all(not image['target_type'] for image in images)
assert '林晚' in next(image['source_context'] for image in images if image['source_index']==1)
assert imported['text'].index('古铜钥匙') < imported['text'].index('第1章')
repeat=client.post('/api/agent/documents/extract',content=word,headers={**headers,'X-Work-Id':str(wid),'X-File-Name':'test.docx'}).json()
assert repeat['document_id']==doc_id and {i['id'] for i in repeat['images']}=={i['id'] for i in images}
assert client.post('/api/agent/documents/extract',content=word,headers={**headers,'X-File-Name':'test.docx'}).status_code==400
header_doc=Document();header_doc.add_paragraph('组织标识参考')
header_doc.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(png('yellow')))
header_buffer=io.BytesIO();header_doc.save(header_buffer)
barrier=threading.Barrier(2)
original_parse=document_imports.parse_word
def concurrent_parse(data):
    barrier.wait(timeout=15)
    return original_parse(data)
document_imports.parse_word=concurrent_parse
before_files=len(list((root/'images').rglob('*.png')))
try:
    with ThreadPoolExecutor(max_workers=2) as pool:
        duplicates=list(pool.map(lambda _: client.post('/api/agent/documents/extract',content=header_buffer.getvalue(),headers={**headers,'X-Work-Id':str(other_work),'X-File-Name':'header.docx'}),range(2)))
finally:
    document_imports.parse_word=original_parse
assert all(result.status_code==200 for result in duplicates)
assert duplicates[0].json()['document_id']==duplicates[1].json()['document_id']
assert '未定位' in duplicates[0].json()['images'][0]['source_context']
assert len(list((root/'images').rglob('*.png')))==before_files+1
assert client.get(f'/api/works/{wid}/resource-images',headers=other_headers).status_code==404
first=next(image for image in images if image['source_index']==1)
second=next(image for image in images if image['source_index']==2)
assert client.get(f"/api/resource-images/{first['id']}/content",headers=other_headers).status_code==404
assert client.put(f"/api/resource-images/{first['id']}",headers=headers,json={'target_type':'character','target_id':cross['id']}).status_code==400
bound=client.put(f"/api/resource-images/{first['id']}",headers=headers,json={
    'target_type':'character','target_id':entity['id'],'label':'林晚·冬季制服','variant_kind':'outfit'}).json()
assert bound['target_id']==entity['id'] and not bound['selected'] and bound['legacy_image_id']
assert len(db.list_entity_images(entity['id'],uid))==1
assert not db.get_entity_image_record(entity['id'],uid)['image_path']
assert client.post(f"/api/resource-images/{first['id']}/select",headers=headers,json={}).status_code==200
assert db.get_entity_image_record(entity['id'],uid)['image_path']
assert client.put(f"/api/resource-images/{first['id']}",headers=headers,json={'target_type':'card','target_id':place['id']}).status_code==400
assert client.put(f"/api/resource-images/{second['id']}",headers=headers,json={'target_type':'card','target_id':place['id'],'label':'雾港·完整广场'}).status_code==200
copy=client.post(f"/api/resource-images/{second['id']}/copy",headers=headers,json={}).json()
assert not copy['target_type']
assert client.put(f"/api/resource-images/{copy['id']}",headers=headers,json={'target_type':'card','target_id':place['id'],'variant_kind':'damaged','parent_image_id':second['id'],'label':'雾港·灾后广场'}).status_code==200
assert client.put(f"/api/resource-images/{second['id']}",headers=headers,json={'parent_image_id':copy['id']}).status_code==400
assert client.put(f"/api/resource-images/{copy['id']}",headers=headers,json={'parent_image_id':first['id']}).status_code==400
for target in (prop,org):
    uploaded=client.post(f"/api/works/{wid}/resource-images?target_type=card&target_id={target['id']}",content=png('blue'),headers={**headers,'X-File-Name':'reference.png'})
    assert uploaded.status_code==200,uploaded.text
assert client.post(f'/api/works/{wid}/resource-images',content=b'not an image',headers=headers).status_code==400

cfg={'work_id':wid}
read=main._tool_read_imported_document(uid,None,cfg,{'document_id':doc_id,'limit':1})
assert read['next_offset']==1 and len(read['prose_sections'])==2
assert main._tool_read_imported_document(uid,None,{'work_id':other_work},{'document_id':doc_id}).get('error')
assert main._tool_visual_asset(uid,None,{'work_id':other_work},{'image_id':first['id']},'inspect').get('error')
view=main._tool_visual_asset(uid,None,cfg,{'image_id':first['id']},'inspect')
assert view['_pi_images'][0]['mimeType']=='image/jpeg'
context=main._agent_documents_system([{'name':'test.docx','text':imported['text'],'document_id':doc_id}],uid,wid)
assert '已导入文档编号' in context['content'] and str(first['id']) in context['content']
chapter=main._tool_import_document_chapter(uid,None,cfg,{'document_id':doc_id,'section_index':1})
assert chapter['changed']
body=db.get_chapter(chapter['chapter_id'],uid)['content']
assert '这是已经写好的正文' in body and '未来计划' not in body
again=main._tool_import_document_chapter(uid,None,cfg,{'document_id':doc_id,'section_index':1})
assert again['already_imported'] and again['chapter_id']==chapter['chapter_id']
assert client.put(f"/api/resource-images/{copy['id']}",headers=headers,json={'chapter_id':chapter['chapter_id']}).status_code==200
plan=db.create_story_plan_node(wid,uid,{'node_type':'book','title':'未来故事总纲','summary':'下一阶段港口毁坏'})
export=client.get(f'/api/works/{wid}/creative-package',headers=headers)
assert export.status_code==200,export.text
with zipfile.ZipFile(io.BytesIO(export.content)) as archive:
    manifest=json.loads(archive.read('manifest.json'))
    assert manifest['format']=='writehtml-creative-package' and manifest['story_plans'][0]['id']==plan['id']
    assets={image['id']:image for image in manifest['images']}
    assert assets[first['id']]['target_ref']==f"character:{entity['id']}"
    assert assets[copy['id']]['target_ref']==f"card:{place['id']}"
    assert assets[copy['id']]['parent_image_ref']==f"image:{second['id']}"
    assert assets[copy['id']]['chapter_ref']==f"chapter:{chapter['chapter_id']}"
    assert 'not_happened_facts' in manifest['semantics']['story_plans']
    for image in assets.values():
        assert hashlib.sha256(archive.read(image['file'])).hexdigest()==image['sha256']
    assert archive.read(f"prose/chapter-{chapter['chapter_id']}.md").decode()==body
assert client.delete(f"/api/resource-images/{second['id']}",headers=headers).status_code==200
assert visual_assets.get_image(copy['id'],uid)['parent_image_id'] is None
assert db.archive_production_card(place['id'],uid)
assert db.delete_chapter(chapter['chapter_id'],uid)
archived_export=client.get(f'/api/works/{wid}/creative-package',headers=headers)
assert archived_export.status_code==200,archived_export.text
with zipfile.ZipFile(io.BytesIO(archived_export.content)) as archive:
    manifest=json.loads(archive.read('manifest.json'))
    archived_card=next(item for item in manifest['cards'] if item['id']==place['id'])
    assert archived_card['status']=='archived'
    archived_image=next(item for item in manifest['images'] if item['id']==copy['id'])
    assert archived_image['target_name']=='雾港' and archived_image['target_ref']==f"card:{place['id']}"
    assert archived_image['chapter_status']=='deleted'
    assert any(item['id']==chapter['chapter_id'] for item in manifest['deleted_chapters'])
    assert main._tool_import_document_chapter(uid,None,cfg,{'document_id':doc_id,'section_index':1}).get('error')
assert any(skill['name']=='document-visual-import' for skill in db.list_agent_skills(uid))
db.init_db()
assert len(visual_assets.list_images(wid,uid))==4
assert db.delete_work(wid,uid)
with db.get_conn() as conn:
    assert not conn.execute('SELECT 1 FROM resource_images WHERE work_id=?',(wid,)).fetchone()
    assert not conn.execute('SELECT 1 FROM document_imports WHERE work_id=?',(wid,)).fetchone()
print('Word/image/variant/ownership/export regressions passed.')
